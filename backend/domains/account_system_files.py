"""Binary system endpoints and durable attendance export consumer."""
import base64
import io
import json
import re
import secrets
import zipfile
from datetime import datetime, timedelta
from urllib.parse import parse_qs, urlparse, quote
from sqlalchemy import select
from starlette.responses import FileResponse
from .. import auth, db, jobs
from ..context import RequestContext, utcnow
from ..errors import ApiError, success
from ..storage import atomic_write, confined_path
from .account_common import *
from .account_checkin import record_times, time_sections
from .account_meetings import bot_get
from .dialog_core import ensure_user_dialog, send_message
from .upload import generic_upload


def uploaded(c,ctx,data,settings,image):
    a=actor(c,ctx); encoded=text(data,'image64'); name=text(data,'filename'); upload=data.get('image' if image else 'files')
    if encoded:
        match=re.fullmatch(r'data:\s*image/(\w+);base64,(.*)',encoded,re.S)
        if not match: raise ApiError('图片数据无效')
        try: raw=base64.b64decode(match[2],validate=True)
        except ValueError as exc: raise ApiError('图片数据无效') from exc
        if len(raw)>64*1024*1024: raise ApiError('文件大小超限')
        stream=io.BytesIO(raw); name=name or 'image.'+match[1]
    else:
        if not upload or not hasattr(upload,'file'): raise ApiError('请选择文件')
        stream=upload.file; name=name or upload.filename
    result=generic_upload(c,ctx,stream,name,data,settings,image=image or bool(encoded))
    if image:
        key=parse_qs(urlparse(result['path']).query)['key'][0]
        db.insert_row(c,'account_picture_uploads',userid=a.userid,month=utcnow().strftime('%Y%m'),name=result['name'],download_key=key)
    return success(result)


def imgupload(c,ctx,data,settings): return uploaded(c,ctx,data,settings,True)
def fileupload(c,ctx,data,settings): return uploaded(c,ctx,data,settings,False)


def imgview(c,ctx,data,settings):
    a=actor(c,ctx); path=text(data,'path').replace('|','/').strip('/'); prefix=f'uploads/user/picture/{a.userid}/'
    path=path.removeprefix(prefix).strip('/')
    if path and not re.fullmatch(r'\d{6}',path): raise ApiError('路径错误')
    t=db.tables['account_picture_uploads']; query=select(t).where(t.c.userid==a.userid)
    dirs=[]; files=[]
    if path:
        query=query.where(t.c.month==path)
        dirs.append({'type':'dir','title':'...','path':'','url':'','thumb':url(ctx,'images/other/dir.png'),'inode':0})
        for item in rows(c,query.order_by(t.c.created_at.desc())):
            value=url(ctx,'api/file/bytes?key='+item['download_key']); files.append({'type':'file','title':item['name'],'path':value,'url':value,'thumb':value,'inode':int(item['created_at'].timestamp())})
    else:
        for month in c.execute(select(t.c.month).where(t.c.userid==a.userid).distinct().order_by(t.c.month.desc())).scalars():
            dirs.append({'type':'dir','title':month,'path':prefix+month,'url':url(ctx,prefix+month),'thumb':url(ctx,'images/other/dir.png'),'inode':int(datetime.strptime(month,'%Y%m').timestamp())})
    return success({'dirs':dirs,'files':files})


def checkin_export(c,ctx,data,settings):
    a=actor(c,ctx,True); config=db.setting(c,'checkinSetting')
    if config.get('open')!='open': raise ApiError('此功能未开启，请前往系统设置开启')
    users=ids(data.get('userid')); dates=data.get('date'); times=data.get('time')
    if not users or not dates or not times: raise ApiError('参数错误')
    if len(users)>100: raise ApiError('导出成员限制最多100个')
    try:
        if not isinstance(dates,list) or len(dates)!=2: raise ValueError()
        start,end=(datetime.strptime(x,'%Y-%m-%d') for x in dates)
    except (TypeError,ValueError) as exc: raise ApiError('日期选择错误') from exc
    if (end-start).days>35: raise ApiError('日期范围限制最大35天')
    if end<start: raise ApiError('日期选择错误')
    if not isinstance(times,list) or len(times)!=2 or not all(re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d',str(x)) for x in times): raise ApiError('时间选择错误')
    bot=bot_get(c,'system-msg'); dialog=ensure_user_dialog(c,bot['userid'],a.userid)
    jobs.enqueue(c,'accounts.checkin_export',{'userid':a.userid,'userids':users,'date':dates,'time':times,'delay':integer(config.get('delay')) or 120,'dialog_id':dialog,'bot_id':bot['userid'],'base_url':ctx.base_url,'language':ctx.language})
    send_message(c,dialog_id=dialog,sender_id=bot['userid'],msg_type='template',msg={'type':'content','content':'正在导出签到数据，请稍等...'},context=ctx)
    return success()


def export_job(c,payload,settings):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    operator=auth.lock_user(c,payload['userid'])
    if auth.is_disabled(operator) or 'admin' not in auth.split_values(operator['identity']):
        raise ApiError('导出发起管理员已失去权限')
    ctx=RequestContext(request_id='job:'+str(payload['_event_id']),base_url=payload['base_url'],language=payload['language'])
    start,end=(datetime.strptime(x,'%Y-%m-%d') for x in payload['date']); times=payload['time']; minute=lambda x:int(x[:2])*60+int(x[3:5]); firstminute,lastminute=map(minute,times)
    book=Workbook(); book.remove(book.active); count=0; first_name=''
    for userid in payload['userids']:
        user=row(c,'users',userid=userid)
        if not user: continue
        count+=1; name=user['nickname'] or str(userid); first_name=first_name or name
        sheet=book.create_sheet(re.sub(r'[\\/*?:\[\]]','',f'{count}.{name}')[:31]); sheet.append(['签到人','签到日期','班次时间','首次签到时间','首次签到结果','最后签到时间','最后签到结果','参数数据'])
        for cell in sheet[1]: cell.font=Font(bold=True)
        records=record_times(c,userid,start,end+timedelta(days=1,seconds=-1)); day=start
        while day<=end:
            values=records.get(day.strftime('%Y-%m-%d'),[])
            actual=[day+timedelta(minutes=minute(v),days=int(minute(v)<firstminute)) for v in values]
            first=next((v for v in actual if day<=v<=day+timedelta(minutes=lastminute,seconds=-1)),None)
            candidates=[v for v in actual if day+timedelta(minutes=firstminute,seconds=1)<=v<=day+timedelta(days=1,minutes=payload['delay'])]; last=candidates[-1] if candidates else None
            first_result='-' if utcnow()<day+timedelta(minutes=firstminute) else '缺卡' if first is None else '迟到' if first>day+timedelta(minutes=firstminute) else '正常'
            last_result='-' if utcnow()<day+timedelta(minutes=lastminute) else '缺卡' if last is None or last==first else '早退' if last<day+timedelta(minutes=lastminute) else '正常'
            if last_result=='-': last=None
            sheet.append([f'{name} (ID: {userid})',day.strftime('%Y-%m-%d'),'-'.join(times),first.strftime('%H:%M') if first else '-',first_result,last.strftime('%H:%M') if last else '-',last_result,', '.join(a+'-'+(b or 'None') for a,b in time_sections(values))])
            for column,result in ((5,first_result),(7,last_result)):
                if result in ('缺卡','迟到','早退'): sheet.cell(sheet.max_row,column).font=Font(color='FF0000' if result=='缺卡' else '436FF6')
            day+=timedelta(days=1)
    if not count:
        send_message(c,dialog_id=payload['dialog_id'],sender_id=payload['bot_id'],msg_type='template',msg={'type':'content','title':'导出签到数据已完成','content':[{'content':'导出签到数据已完成','style':'font-weight: bold;padding-bottom: 4px;'}, {'content':'没有任何数据','style':'color: #ff0000;'}]},context=ctx)
        return
    name=first_name+(f'等{count}位成员' if count>1 else '')+'的签到记录_'+str(payload['_event_id'])+'.xlsx'; name=re.sub(r'[\\/:*?"<>|]','_',name)
    output=io.BytesIO(); book.save(output); book.close(); archive=io.BytesIO()
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as zipped: zipped.writestr(name,output.getvalue())
    relative='temp/checkin/export/'+str(payload['_event_id'])+'.zip'; atomic_write(settings.storage_root,relative,archive.getvalue())
    key=secrets.token_hex(32)
    db.insert_row(c,'backend_file_downloads',id=key,userid=payload['userid'],path=relative,name=name.removesuffix('.xlsx')+'.zip',ready=1,expires_at=utcnow()+timedelta(days=1))
    send_message(c,dialog_id=payload['dialog_id'],sender_id=payload['bot_id'],msg_type='template',msg={'type':'file_download','title':'导出签到数据已完成','name':name,'size':len(archive.getvalue()),'url':url(ctx,'api/system/checkin/down?key='+key)},context=ctx)


def checkin_down(c,ctx,data,settings):
    key=text(data,'key'); value=row(c,'backend_file_downloads',id=key)
    if not value or not value['ready'] or value['expires_at']<=utcnow() or not value['path'].startswith('temp/checkin/export/'):
        from starlette.responses import JSONResponse
        return JSONResponse({'ret':0,'msg':'文件不存在！','data':[]},status_code=403)
    users=db.tables['users']
    owner=c.execute(select(users).where(users.c.userid==value['userid']).with_for_update()).mappings().first()
    if not owner or auth.is_disabled(owner) or 'admin' not in auth.split_values(owner['identity']):
        from starlette.responses import JSONResponse
        return JSONResponse({'ret':0,'msg':'文件不存在！','data':[]},status_code=403)
    path=confined_path(settings.storage_root,value['path'])
    if not path.is_file(): raise ApiError('文件不存在！')
    return FileResponse(path,filename=value['name'],media_type='application/zip')




def register_jobs(): jobs.register('accounts.checkin_export',export_job)

ROUTES={'system/imgupload':imgupload,'system/imgview':imgview,'system/fileupload':fileupload,'system/checkin/export':checkin_export,'system/checkin/down':checkin_down}
