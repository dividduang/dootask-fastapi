"""Administrator-controlled account creation and spreadsheet import."""
import csv
import io
from pathlib import Path
from starlette.responses import Response
from sqlalchemy import select
from .. import auth, db
from ..errors import ApiError, success
from .account_common import *
from .account_admin import checked_departments


def validate_import(c,value):
    if not isinstance(value,dict): raise ApiError('参数错误')
    email,nickname,password=(text(value,k) for k in ('email','nickname','password'))
    if not email or not nickname or not password: raise ApiError('邮箱、昵称、初始密码均为必填')
    if not email_valid(email): raise ApiError('邮箱格式不正确')
    if not 2<=len(nickname)<=20: raise ApiError('昵称需为2-20个字')
    auth.validate_password(c,password); named(text(value,'profession'),'职位/职称')


def createuser(c,ctx,data,settings):
    actor(c,ctx,True); validate_import(c,data)
    departments=checked_departments(c,data.get('department',[]))
    user=auth.create_account(c,ctx,text(data,'email'),text(data,'password'),nickname=text(data,'nickname'),
        profession=text(data,'profession'),departments=departments,
        changepass=integer(data.get('changepass'),1)==1,email_verity=integer(data.get('email_verity'))==1)
    return success(safe_user(c,user,ctx),'创建成功')


def parse_sheet(sheet):
    result=[]
    for index,cells in enumerate(sheet):
        if index==0: continue
        values=[str(x).strip() if x is not None else '' for x in list(cells)[:4]]
        values+=['']*(4-len(values))
        if not any(values[:3]): continue
        result.append(dict(zip(('email','nickname','password','profession'),values),line=index+1))
    return result


def import_preview(c,ctx,data,settings):
    actor(c,ctx,True); upload=data.get('file')
    if not upload or not hasattr(upload,'file'): raise ApiError('请选择文件')
    ext=Path(upload.filename or '').suffix.lower()
    if ext not in ('.xls','.xlsx','.csv'): raise ApiError('仅支持 xls/xlsx/csv 文件')
    if ext=='.csv':
        raw=upload.file.read(16*1024*1024+1)
        if len(raw)>16*1024*1024: raise ApiError('文件过大')
        try: sheet=list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
        except UnicodeDecodeError: sheet=list(csv.reader(io.StringIO(raw.decode('gb18030'))))
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(upload.file,read_only=True,data_only=True)
        try: sheet=list(book.worksheets[0].iter_rows(values_only=True))
        finally: book.close()
    else:
        import xlrd
        book=xlrd.open_workbook(file_contents=upload.file.read()); first=book.sheet_by_index(0)
        sheet=[first.row_values(i) for i in range(first.nrows)]
    values=parse_sheet(sheet)
    if not values: raise ApiError('文件中没有可导入的数据')
    if len(values)>500: raise ApiError('单次最多导入500条')
    t=db.tables['users']; existing=set(str(v).lower() for v in c.execute(select(t.c.email).where(t.c.email.in_([v['email'] for v in values]))).scalars()); seen=set(); valid=0
    for value in values:
        reason=''
        try:
            validate_import(c,value); email=value['email'].lower()
            if email in seen: raise ApiError('文件内邮箱重复')
            seen.add(email)
            if email in existing: raise ApiError('邮箱地址已存在')
        except ApiError as exc: reason=exc.message
        value.update(email_verity=1,status='error' if reason else 'ok',reason=reason)
        valid+=not bool(reason)
    return success({'total':len(values),'valid':valid,'invalid':len(values)-valid,'rows':values},'解析完成')


def import_users(c,ctx,data,settings):
    actor(c,ctx,True); values=data.get('rows')
    if not isinstance(values,list) or not values: raise ApiError('没有可导入的数据')
    if len(values)>500: raise ApiError('单次最多导入500条')
    failed=[]; seen=set(); count=0
    for value in values:
        try:
            validate_import(c,value); email=text(value,'email').lower()
            if email in seen: raise ApiError('文件内邮箱重复')
            seen.add(email)
            with c.begin_nested(): createuser(c,ctx,dict(value,changepass=data.get('changepass',1)),settings)
            count+=1
        except ApiError as exc:
            failed.append({'line':value.get('line',0) if isinstance(value,dict) else 0,'email':value.get('email','') if isinstance(value,dict) else '', 'reason':exc.message})
    return success({'total':len(values),'success':count,'failed':failed},'导入完成')


def import_template(c,ctx,data,settings):
    actor(c,ctx,True)
    from openpyxl import Workbook
    book=Workbook(); sheet=book.active
    sheet.append(['邮箱(必填)','昵称(必填,2-20字)','初始密码(必填,6-32位)','职位(选填,2-20字)'])
    sheet.append(['employee@example.com','张三','Abc123456','工程师'])
    output=io.BytesIO(); book.save(output); book.close()
    return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="user_import_template.xlsx"'})


ROUTES={'users/createuser':createuser,'users/import/preview':import_preview,'users/import':import_users,'users/import/template':import_template}
